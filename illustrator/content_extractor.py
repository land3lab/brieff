"""Turn an input source (pptx / pdf / docx / txt / raw text) into a list of
short content chunks, one per slide/section, ready to be summarized into an
image concept.
"""
from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from PIL import Image


@dataclass
class ContentChunk:
    index: int
    title: str
    body: str

    @property
    def combined_text(self) -> str:
        return f"{self.title}\n{self.body}".strip()


@dataclass
class ImageChunk:
    """A slide's content given as a pasted/uploaded screenshot rather than text.

    The image is sent to a vision-capable chat model to derive the visual
    concept, so no local OCR dependency is needed.
    """

    index: int
    image: Image.Image
    title: str = "(붙여넣은 이미지)"


def _chunk_from_texts(texts: list[tuple[str, str]]) -> list[ContentChunk]:
    chunks = []
    for i, (title, body) in enumerate(texts):
        if not title.strip() and not body.strip():
            continue
        chunks.append(ContentChunk(index=i, title=title.strip(), body=body.strip()))
    return chunks


def extract_from_pptx(path: str | Path) -> list[ContentChunk]:
    from pptx import Presentation

    prs = Presentation(str(path))
    texts = []
    for slide in prs.slides:
        title = ""
        body_lines = []
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            text = shape.text_frame.text.strip()
            if not text:
                continue
            is_title = shape == slide.shapes.title
            if is_title and not title:
                title = text
            else:
                body_lines.append(text)
        texts.append((title, "\n".join(body_lines)))
    return _chunk_from_texts(texts)


def extract_from_pdf(path: str | Path) -> list[ContentChunk]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    texts = []
    for page in reader.pages:
        raw = (page.extract_text() or "").strip()
        if not raw:
            texts.append(("", ""))
            continue
        lines = [l for l in raw.splitlines() if l.strip()]
        title = lines[0] if lines else ""
        body = "\n".join(lines[1:])
        texts.append((title, body))
    return _chunk_from_texts(texts)


def extract_from_docx(path: str | Path) -> list[ContentChunk]:
    import docx

    document = docx.Document(str(path))
    sections: list[tuple[str, list[str]]] = []
    for para in document.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        is_heading = para.style.name.lower().startswith("heading") or para.style.name.lower() == "title"
        if is_heading or not sections:
            sections.append((text if is_heading else "", [] if is_heading else [text]))
        else:
            sections[-1][1].append(text)
    texts = [(title, "\n".join(body)) for title, body in sections]
    return _chunk_from_texts(texts)


def extract_from_text(raw_text: str) -> list[ContentChunk]:
    blocks = [b.strip() for b in raw_text.split("\n\n") if b.strip()]
    if not blocks:
        return []
    texts = []
    for block in blocks:
        lines = block.splitlines()
        title = lines[0] if len(lines) > 1 else ""
        body = "\n".join(lines[1:]) if len(lines) > 1 else lines[0]
        texts.append((title, body))
    return _chunk_from_texts(texts)


def extract_from_file(path: str | Path) -> list[ContentChunk]:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".pptx":
        return extract_from_pptx(path)
    if suffix == ".pdf":
        return extract_from_pdf(path)
    if suffix == ".docx":
        return extract_from_docx(path)
    if suffix in (".txt", ".md"):
        return extract_from_text(path.read_text(encoding="utf-8"))
    raise ValueError(f"지원하지 않는 파일 형식입니다: {suffix}")
