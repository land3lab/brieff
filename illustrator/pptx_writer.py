"""Insert generated illustrations back into a copy of the source .pptx."""
from __future__ import annotations

import io
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.util import Inches


def embed_images(
    source_pptx: str | Path,
    images_by_slide_index: dict[int, Image.Image],
    output_path: str | Path,
    width_inches: float = 3.0,
    margin_inches: float = 0.3,
) -> Path:
    prs = Presentation(str(source_pptx))
    height_inches = width_inches * 3 / 4

    for slide_index, slide in enumerate(prs.slides):
        image = images_by_slide_index.get(slide_index)
        if image is None:
            continue
        buffer = io.BytesIO()
        image.save(buffer, format="PNG")
        buffer.seek(0)

        left = prs.slide_width - Inches(width_inches) - Inches(margin_inches)
        top = prs.slide_height - Inches(height_inches) - Inches(margin_inches)
        slide.shapes.add_picture(
            buffer, left, top, width=Inches(width_inches), height=Inches(height_inches)
        )

    output_path = Path(output_path)
    prs.save(str(output_path))
    return output_path
